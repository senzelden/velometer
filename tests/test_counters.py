"""M8 Task 6: counter framework (hourly to daily, chained index, check) and the Berlin, Paris and
Vienna adapters on fixtures cut from the real portal files (tests/fixtures/counters/<city>/).

Day totals below were summed by hand from the fixture rows (docs/review/m8-counters.md).
"""

import copy
import json
import shutil
import sqlite3
import zipfile
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
import pytest
import yaml
from typer.testing import CliRunner

from velometer.cli import app
from velometer.counters.adapters import berlin, paris, vienna
from velometer.counters.adapters import download as download_files
from velometer.counters.check import check_counters
from velometer.counters.days import sum_counters, to_days, zero_days_invalid
from velometer.counters.index import MIN_SITES, chain_index, chain_index_monthly
from velometer.counters.model import (
    CounterEntry,
    CounterFormatError,
    CountersConfig,
    DayCount,
    Grain,
    HourCount,
    MonthValue,
    load_counters_config,
)
from velometer.db import (
    CounterIndexYear,
    CounterSite,
    create_schema,
    replace_counter_index,
    replace_counter_sites,
)

FIX = Path(__file__).parent / "fixtures" / "counters"
CONFIG = Path(__file__).resolve().parents[1] / "config"
BERLIN_XLSX = FIX / "berlin" / "gesamtdatei-stundenwerte.xlsx"
PARIS_FILES = [FIX / "paris" / f for f in ("2019.zip", "2020.zip", "2022.zip", "2025.zip")] + [
    FIX / "paris" / "rolling.csv"
]
VIENNA_CSV = FIX / "vienna" / "radverkehrszaehlungen.csv"


def _entry(city: str, adapter: str, tz: str, grain: Grain = "hourly", **kw) -> CounterEntry:
    return CounterEntry(
        city=city,
        adapter=adapter,
        source="test",
        licence="test-licence",
        licence_url="https://example.org/licence",
        page="https://example.org/page",
        urls=["https://example.org/a.csv"],
        timezone=tz,
        grain=grain,
        **kw,
    )


def _by_key(days):
    return {(d.site_id, d.day): d for d in days}


# --- hourly to daily ------------------------------------------------------------------------


def _hours(site, day, hours, count=1, fold=0):
    return [
        HourCount(site, datetime(day.year, day.month, day.day, h, fold=fold), count) for h in hours
    ]


def test_spring_forward_day_with_23_hours_is_complete():
    day = date(2023, 3, 26)  # last Sunday of March: 02:00 does not exist in Europe/Berlin
    hours = _hours("a", day, [h for h in range(24) if h != 2], count=2)
    (d,) = to_days(hours, "Europe/Berlin")
    assert (d.day, d.count, d.complete) == (day, 46, True)


def test_spring_forward_placeholder_hour_is_ignored():
    day = date(2023, 3, 26)
    hours = _hours("a", day, [h for h in range(24) if h != 2]) + _hours("a", day, [2], count=50)
    (d,) = to_days(hours, "Europe/Berlin")
    assert (d.count, d.complete) == (23, True)


def test_fall_back_day_needs_25_hours():
    day = date(2023, 10, 29)  # last Sunday of October: 02:00 occurs twice
    first = _hours("a", day, range(24))
    second = _hours("a", day, [2], count=5, fold=1)
    (d25,) = to_days(first + second, "Europe/Berlin")
    assert (d25.count, d25.complete) == (29, True)
    (d24,) = to_days(first, "Europe/Berlin")
    assert (d24.count, d24.complete) == (24, False)


def test_ordinary_day_needs_24_hours():
    day = date(2023, 6, 1)
    assert to_days(_hours("a", day, range(24)), "Europe/Berlin")[0].complete
    assert not to_days(_hours("a", day, range(23)), "Europe/Berlin")[0].complete


def test_duplicate_hours_from_overlapping_files_count_once_last_wins():
    day = date(2023, 6, 1)
    file_a = _hours("a", day, range(24), count=1)
    file_b = _hours("a", day, range(12, 24), count=3)  # overlaps the afternoon of file_a
    (d,) = to_days(file_a + file_b, "Europe/Berlin")
    assert (d.count, d.complete) == (12 * 1 + 12 * 3, True)


def test_zero_day_is_an_outage():
    days = [DayCount("a", date(2024, 1, 1), 0, True), DayCount("a", date(2024, 1, 2), 5, True)]
    assert [d.complete for d in zero_days_invalid(days)] == [False, True]


# --- chained index --------------------------------------------------------------------------


def _year(site, year, per_day, share=1.0):
    start = date(year, 1, 1)
    n_days = (date(year + 1, 1, 1) - start).days
    n_complete = round(n_days * share)
    return [
        DayCount(site, start + timedelta(days=i), per_day, i < n_complete) for i in range(n_days)
    ]


def test_base_year_is_100():
    days = [d for s in "abc" for d in _year(s, 2022, 10)]
    assert chain_index(days, 2022) == [(2022, 100.0, 3)]


def test_index_continues_on_remaining_sites_after_a_site_drops_out():
    days = []
    for y, per_day in ((2021, 10), (2022, 20), (2023, 30)):
        for s in "abc":
            days += _year(s, y, per_day)
    days += _year("d", 2021, 1000) + _year("d", 2022, 1000)  # d stops after year 2 (2022)
    out = {y: (idx, n) for y, idx, n in chain_index(days, 2022)}
    assert out[2022] == (100.0, 4)
    assert out[2021] == (pytest.approx(100.0 * (3 * 10 + 1000) / (3 * 20 + 1000)), 4)
    assert out[2023] == (pytest.approx(100.0 * 30 / 20), 3)


def test_two_qualifying_sites_give_no_value_and_the_chain_stays_broken():
    days = []
    for y in (2021, 2022, 2023, 2024, 2025):
        for s in "abc":
            if not (y == 2023 and s == "c"):  # 2023 has only a and b
                days += _year(s, y, 10)
    out = {y: (idx, n) for y, idx, n in chain_index(days, 2021)}
    assert out[2022] == (pytest.approx(100.0), 3)
    assert out[2023] == (None, 2)
    assert out[2024] == (None, 2)
    # 2024 -> 2025 has three qualifying sites, but the chain is not re-based after a break
    assert out[2025] == (None, 3)


def test_a_break_before_the_base_year_stays_broken_backwards():
    days = []
    for y in (2021, 2022, 2023, 2024, 2025):
        for s in "abc":
            if not (y == 2023 and s == "c"):  # 2023 has only a and b
                days += _year(s, y, 10)
    out = {y: (idx, n) for y, idx, n in chain_index(days, 2025)}
    assert out[2024] == (pytest.approx(100.0 * 366 / 365), 3)  # 2024 has 366 days of 10
    assert out[2023] == (None, 2)
    assert out[2022] == (None, 2)
    # 2021 -> 2022 has three qualifying sites, but the chain is not re-based before the base
    assert out[2021] == (None, 3)


def test_a_site_below_90_percent_complete_days_does_not_qualify():
    days = []
    for s in "abc":
        days += _year(s, 2022, 10) + _year(s, 2023, 10)
    days += _year("d", 2022, 10) + _year("d", 2023, 99, share=0.89)
    out = {y: n for y, _, n in chain_index(days, 2022)}
    assert out[2023] == 3


# --- adapters -------------------------------------------------------------------------------


def test_berlin_fixture_sites_and_day_totals():
    entry = _entry("berlin", "berlin", "Europe/Berlin")
    sites, hours = berlin.parse([BERLIN_XLSX], entry)
    by_id = {s.site_id: s for s in sites}
    assert set(by_id) == {"02-MI-JAN-N", "05-FK-OBB-O", "12-PA-SCH"}
    assert by_id["02-MI-JAN-N"].name == "Jannowitzbrücke Nord"
    assert by_id["02-MI-JAN-N"].lat == pytest.approx(52.513932, abs=1e-6)
    days = _by_key(to_days(hours, entry.timezone))
    jan = "02-MI-JAN-N"
    assert (days[jan, date(2018, 12, 31)].count, days[jan, date(2018, 12, 31)].complete) == (
        887,
        True,
    )
    assert days[jan, date(2019, 1, 1)].count == 458
    assert (days[jan, date(2019, 3, 31)].count, days[jan, date(2019, 3, 31)].complete) == (
        1924,
        True,
    )
    # fall-back day: the file has one 02:00 row, so 24 of 25 hours
    assert (days[jan, date(2019, 10, 27)].count, days[jan, date(2019, 10, 27)].complete) == (
        1643,
        False,
    )
    # every Oberbaumbrücke cell on 2019-10-27 carries the "no or implausible data" fill
    assert ("05-FK-OBB-O", date(2019, 10, 27)) not in days
    assert days["05-FK-OBB-O", date(2019, 1, 1)].count == 813


def _berlin_copy(tmp_path, edit) -> Path:
    path = tmp_path / "berlin.xlsx"
    wb = openpyxl.load_workbook(BERLIN_XLSX)
    edit(wb)
    wb.save(path)
    return path


def test_berlin_no_data_fill_on_a_positive_value_is_a_missing_hour(tmp_path):
    """The legend's second fill ("keine Zähldaten ... vorhanden", FFFFEB9C) marks missing data
    like the orange one; the real file has it only on empty cells and seven 0s (2022), so the
    fixture copy puts it on a positive value."""

    def flag(wb):
        legend_fill = wb["Legende"]["B3"].fill
        assert legend_fill.fgColor.rgb == "FFFFEB9C"
        ws = wb["Jahresdatei 2019"]
        for row in ws.iter_rows(min_row=2):
            if row[0].value == datetime(2019, 1, 1, 12):
                assert row[1].value > 0  # 02-MI-JAN-N
                row[1].fill = copy.copy(legend_fill)

    entry = _entry("berlin", "berlin", "Europe/Berlin")
    _, hours = berlin.parse([_berlin_copy(tmp_path, flag)], entry)
    _, real = berlin.parse([BERLIN_XLSX], entry)
    noon = next(
        h.count
        for h in real
        if h.site_id == "02-MI-JAN-N" and h.local_hour == datetime(2019, 1, 1, 12)
    )
    day = _by_key(to_days(hours, entry.timezone))["02-MI-JAN-N", date(2019, 1, 1)]
    assert (day.count, day.complete) == (458 - noon, False)


def test_berlin_legend_without_a_missing_data_row_raises(tmp_path):
    def drop(wb):
        for row in wb["Legende"].iter_rows():
            for c in row:
                if isinstance(c.value, str) and "keine" in c.value:
                    c.value = "x"

    with pytest.raises(CounterFormatError, match="Legende"):
        berlin.parse([_berlin_copy(tmp_path, drop)], _entry("berlin", "berlin", "Europe/Berlin"))


def test_berlin_renamed_column_raises_naming_it(tmp_path):
    def rename(wb):
        ws = wb["Standortdaten"]
        for c in ws[1]:
            if c.value == "Breitengrad":
                c.value = "Breite"

    path = _berlin_copy(tmp_path, rename)
    with pytest.raises(CounterFormatError, match="Breitengrad"):
        berlin.parse([path], _entry("berlin", "berlin", "Europe/Berlin"))


def test_berlin_renumbered_header_needs_an_alias(tmp_path):
    def renumber(wb):
        ws = wb["Jahresdatei 2019"]
        for c in ws[1]:
            if c.value and str(c.value).startswith("05-FK-OBB-O"):
                c.value = "09-FK-OBB-O 01.06.2015"

    path = _berlin_copy(tmp_path, renumber)
    with pytest.raises(CounterFormatError, match="09-FK-OBB-O"):
        berlin.parse([path], _entry("berlin", "berlin", "Europe/Berlin"))
    entry = _entry("berlin", "berlin", "Europe/Berlin", site_aliases={"09-FK-OBB-O": "05-FK-OBB-O"})
    _, hours = berlin.parse([path], entry)
    days = _by_key(to_days(hours, entry.timezone))
    assert days["05-FK-OBB-O", date(2019, 1, 1)].count == 813


def test_paris_fixture_sites_and_day_totals():
    entry = _entry("paris", "paris", "Europe/Paris")
    sites, hours = paris.parse(PARIS_FILES, entry)
    by_id = {s.site_id: s for s in sites}
    assert set(by_id) == {"100003096", "100036718"}
    assert by_id["100003096"].name == "97 avenue Denfert Rochereau"
    assert (by_id["100003096"].lat, by_id["100003096"].lon) == (48.83504, 2.33314)
    days = _by_key(zero_days_invalid(to_days(hours, entry.timezone)))
    s1, s2 = "100003096", "100036718"
    # 2019 archive, offsets: 23 hours on the spring-forward day, the placeholder 03:00+02:00
    # row (count 0) is the same hour once
    assert (days[s1, date(2019, 3, 31)].count, days[s1, date(2019, 3, 31)].complete) == (604, True)
    # two direction counters of one site are summed
    assert days[s2, date(2019, 3, 31)].count == 1088
    # the 2019 archive ends at 2019-12-31 00:00
    assert (days[s1, date(2019, 12, 31)].count, days[s1, date(2019, 12, 31)].complete) == (3, False)
    assert (days[s1, date(2020, 1, 1)].count, days[s1, date(2020, 1, 1)].complete) == (234, True)
    assert days[s2, date(2020, 1, 1)].count == 655
    # 2022 archive, naive local times with a 02:00 row on the spring-forward day
    assert (days[s1, date(2022, 3, 27)].count, days[s1, date(2022, 3, 27)].complete) == (953, True)
    assert (days[s2, date(2022, 3, 27)].count, days[s2, date(2022, 3, 27)].complete) == (0, False)
    # 2025-10-25 is in the 2025 archive (cp1252, commas) and in the rolling export (UTC): once
    assert (days[s1, date(2025, 10, 25)].count, days[s1, date(2025, 10, 25)].complete) == (
        989,
        True,
    )
    assert days[s2, date(2025, 10, 25)].count == 2082


def test_paris_renamed_column_raises_naming_it(tmp_path):
    text = (FIX / "paris" / "rolling.csv").read_text(encoding="utf-8-sig")
    bad = tmp_path / "rolling.csv"
    bad.write_text(text.replace("Comptage horaire", "Comptage"), encoding="utf-8")
    with pytest.raises(CounterFormatError, match="Comptage horaire"):
        paris.parse([bad], _entry("paris", "paris", "Europe/Paris"))


def test_paris_reads_a_zip_member_and_a_bare_csv_alike(tmp_path):
    entry = _entry("paris", "paris", "Europe/Paris")
    z = tmp_path / "rolling.zip"
    with zipfile.ZipFile(z, "w") as zf:
        zf.write(FIX / "paris" / "rolling.csv", "rolling.csv")
    assert paris.parse([z], entry)[1] == paris.parse([FIX / "paris" / "rolling.csv"], entry)[1]


def test_paris_skips_the_unlabelled_total_row(tmp_path):
    text = (FIX / "paris" / "rolling.csv").read_text(encoding="utf-8-sig")
    total = ";;;;5000.0;2025-10-25T10:00:00+00:00;null, null\n"
    f = tmp_path / "rolling.csv"
    f.write_text(text + total, encoding="utf-8")
    entry = _entry("paris", "paris", "Europe/Paris")
    sites, hours = paris.parse([f], entry)
    assert {s.site_id for s in sites} == {"100003096", "100036718"}
    assert hours == paris.parse([FIX / "paris" / "rolling.csv"], entry)[1]


def test_paris_a_file_without_coordinates_keeps_the_last_given(tmp_path):
    text = (FIX / "paris" / "rolling.csv").read_text(encoding="utf-8-sig")
    blank = tmp_path / "blank.csv"
    blank.write_text(text.replace("48.83504, 2.33314", ""), encoding="utf-8")
    sites, _ = paris.parse(
        [FIX / "paris" / "rolling.csv", blank], _entry("p", "paris", "Europe/Paris")
    )
    site = next(s for s in sites if s.site_id == "100003096")
    assert (site.lat, site.lon) == (48.83504, 2.33314)


def _rezip(src: Path, dst: Path, keep) -> Path:
    with zipfile.ZipFile(src) as z:
        (name,) = z.namelist()
        lines = z.read(name).splitlines(keepends=True)
    with zipfile.ZipFile(dst, "w") as z:
        z.writestr(name, b"".join([lines[0]] + [ln for ln in lines[1:] if keep(ln)]))
    return dst


def test_paris_a_counter_missing_an_hour_makes_the_site_day_incomplete(tmp_path):
    drop = b"100036718-103036718;", b"2019-03-31T12:00"
    f = _rezip(
        FIX / "paris" / "2019.zip",
        tmp_path / "2019.zip",
        lambda ln: not (ln.startswith(drop[0]) and drop[1] in ln),
    )
    entry = _entry("paris", "paris", "Europe/Paris")
    days = _by_key(to_days(paris.parse([f], entry)[1], entry.timezone))
    assert not days["100036718", date(2019, 3, 31)].complete
    assert days["100003096", date(2019, 3, 31)].complete


def _pc(site, counter, day, hours, count=1):
    return {(site, counter, datetime(day.year, day.month, day.day, h), 0): count for h in hours}


def test_paris_counter_in_service_is_its_reported_span():
    d1, d2, d3 = date(2023, 6, 1), date(2023, 6, 2), date(2023, 6, 3)
    per = {}
    per |= (
        _pc("s", "a", d1, range(24)) | _pc("s", "a", d2, range(24)) | _pc("s", "a", d3, range(24))
    )
    # b reports d1 and d3 but nothing on d2: inside its span, an outage
    per |= _pc("s", "b", d1, range(24)) | _pc("s", "b", d3, range(24))
    # c reports from d2 on: not in service on d1
    per |= _pc("s", "c", d2, range(24)) | _pc("s", "c", d3, range(24))
    days = _by_key(to_days(sum_counters(per), "Europe/Paris"))
    assert (days["s", d1].count, days["s", d1].complete) == (48, True)
    assert ("s", d2) not in days  # no hour has every in-service counter (a, b, c)
    assert (days["s", d3].count, days["s", d3].complete) == (72, True)


def _vienna_entry():
    return _entry("vienna", "vienna", "Europe/Vienna", "daily", exclude_columns=["Opernring"])


def test_vienna_fixture_sites_and_day_totals():
    sites, days = vienna.parse([VIENNA_CSV], _vienna_entry())
    names = {s.site_id for s in sites}
    assert len(names) == 19 and "Opernring" not in names and "Neubaugürtel" in names
    assert all(s.lat is None and s.lon is None for s in sites)
    by = _by_key(days)
    assert by["Donaukanal", date(2023, 12, 31)].count == 395
    assert by["Opernring Innen", date(2013, 1, 9)].count == 1121  # written "1.121"
    assert by["Argentinierstraße", date(2013, 1, 10)].count == 1028
    assert ("Pfeilgasse", date(2013, 1, 9)) not in by  # empty cell: not reported
    assert by["Argentinierstraße", date(2024, 1, 1)].complete  # reported, as 0
    assert not _by_key(zero_days_invalid(days))["Argentinierstraße", date(2024, 1, 1)].complete


def test_vienna_renamed_column_raises_naming_it(tmp_path):
    bad = tmp_path / "v.csv"
    bad.write_text(
        VIENNA_CSV.read_text(encoding="utf-8-sig").replace("Datum", "Tag"), encoding="utf-8"
    )
    with pytest.raises(CounterFormatError, match="Datum"):
        vienna.parse([bad], _vienna_entry())
    bad.write_text(
        VIENNA_CSV.read_text(encoding="utf-8-sig").replace(";Opernring;", ";Oper;"),
        encoding="utf-8",
    )
    with pytest.raises(CounterFormatError, match="Opernring"):
        vienna.parse([bad], _vienna_entry())


# --- fetch, config ----------------------------------------------------------------------------


def test_fetch_writes_files_and_records_url_licence_and_time(tmp_path):
    entry = _entry("x", "vienna", "Europe/Vienna", "daily")
    entry = entry.model_copy(
        update={"urls": ["https://example.org/a/data.csv", "https://example.org/exports/csv?q=1"]}
    )
    paths = download_files(entry, tmp_path, fetch=lambda url: url.encode())
    assert [p.name for p in paths] == ["00_data.csv", "01_csv"]
    meta = json.loads((tmp_path / "retrieved.json").read_text())
    assert meta["licence"] == "test-licence"
    assert [f["url"] for f in meta["files"]] == entry.urls
    assert all(f["retrieved_at"].endswith("Z") for f in meta["files"])


def test_fetch_writes_nothing_when_a_url_fails(tmp_path):
    entry = _entry("x", "vienna", "Europe/Vienna", "daily")
    entry = entry.model_copy(update={"urls": ["https://example.org/a", "https://example.org/b"]})

    def flaky(url):
        if url.endswith("/b"):
            raise OSError("down")
        return b"x"

    with pytest.raises(OSError):
        download_files(entry, tmp_path / "out", fetch=flaky)
    assert not (tmp_path / "out").exists()


def test_counters_config_lists_every_city_once():
    cfg = load_counters_config(CONFIG / "counters.yaml")
    known = {c["id"] for c in yaml.safe_load((CONFIG / "cities.yaml").read_text())["cities"]}
    listed = [c.city for c in cfg.cities] + [n.city for n in cfg.none]
    assert set(listed) <= known and len(listed) == len(set(listed))
    assert len(cfg.none) == 14 and all(n.checked for n in cfg.none)
    adapters = {c.city: c.adapter for c in cfg.cities}
    assert {k: adapters.get(k) for k in ("berlin", "paris", "vienna")} == {
        "berlin": "berlin",
        "paris": "paris",
        "vienna": "vienna",
    }


# --- load (CLI) and check -------------------------------------------------------------------


def _setup_city(tmp_path, city: str, entry: dict, files: list[Path]) -> tuple[Path, Path, Path]:
    cfg_dir = tmp_path / "config"
    cfg_dir.mkdir(exist_ok=True)
    cfg_path = cfg_dir / "counters.yaml"
    cfg = yaml.safe_load(cfg_path.read_text()) if cfg_path.exists() else {"cities": [], "none": []}
    cfg["cities"].append(entry)
    cfg_path.write_text(yaml.safe_dump(cfg, allow_unicode=True), encoding="utf-8")
    out = tmp_path / "data" / "counters" / city
    out.mkdir(parents=True)
    meta_files = []
    for i, f in enumerate(files):
        name = f"{i:02d}_{f.name}"
        shutil.copy(f, out / name)
        meta_files.append(
            {
                "file": name,
                "url": f"https://example.org/{f.name}",
                "retrieved_at": "2026-10-07T12:00:00Z",
            }
        )
    (out / "retrieved.json").write_text(json.dumps({"files": meta_files}), encoding="utf-8")
    return cfg_dir, tmp_path / "data", tmp_path / "v.sqlite"


BERLIN_ENTRY = {
    "city": "berlin",
    "adapter": "berlin",
    "source": "Senatsverwaltung",
    "licence": "dl-de-zero-2.0",
    "licence_url": "https://www.govdata.de/dl-de/zero-2-0",
    "page": "https://example.org/berlin",
    "urls": ["https://example.org/gesamtdatei-stundenwerte.xlsx"],
    "timezone": "Europe/Berlin",
    "grain": "hourly",
}


def _dump(db: Path) -> dict:
    conn = sqlite3.connect(db)
    try:
        return {
            t: conn.execute(f"SELECT * FROM {t} ORDER BY 1, 2, 3").fetchall()
            for t in ("counter_sites", "counter_daily", "counter_index")
        }
    finally:
        conn.close()


def test_load_twice_gives_the_same_rows(tmp_path):
    cfg_dir, data, db = _setup_city(tmp_path, "berlin", BERLIN_ENTRY, [BERLIN_XLSX])
    args = ["counters", "load", "berlin", "--db", str(db), "--data", str(data)]
    args += ["--config-dir", str(cfg_dir)]
    first = CliRunner().invoke(app, args)
    assert first.exit_code == 0, first.output
    rows = _dump(db)
    assert len(rows["counter_sites"]) == 3
    jan = [r for r in rows["counter_daily"] if r[1] == "02-MI-JAN-N" and r[2] == "2019-03-31"]
    assert jan == [("berlin", "02-MI-JAN-N", "2019-03-31", 1924, 1)]
    site = next(r for r in rows["counter_sites"] if r[1] == "02-MI-JAN-N")
    assert site[5:] == (
        "Senatsverwaltung",
        "dl-de-zero-2.0",
        "https://example.org/berlin",
        "2026-10-07T12:00:00Z",
        "2018-12-31",
        "2019-10-27",
    )
    second = CliRunner().invoke(app, args)
    assert second.exit_code == 0, second.output
    assert _dump(db) == rows


def test_load_replaces_days_the_files_no_longer_give(tmp_path):
    cfg_dir, data, db = _setup_city(tmp_path, "berlin", BERLIN_ENTRY, [BERLIN_XLSX])
    args = ["counters", "load", "berlin", "--db", str(db), "--data", str(data)]
    args += ["--config-dir", str(cfg_dir)]
    assert CliRunner().invoke(app, args).exit_code == 0
    rows = _dump(db)
    conn = sqlite3.connect(db)
    with conn:
        conn.execute("INSERT INTO counter_daily VALUES ('berlin', 'gone', '2019-01-01', 5, 1)")
    conn.close()
    assert CliRunner().invoke(app, args).exit_code == 0
    assert _dump(db) == rows


def test_a_failed_replace_keeps_the_previous_daily_rows(tmp_path, monkeypatch):
    import velometer.counters as counters

    cfg_dir, data, db = _setup_city(tmp_path, "berlin", BERLIN_ENTRY, [BERLIN_XLSX])
    args = ["counters", "load", "berlin", "--db", str(db), "--data", str(data)]
    args += ["--config-dir", str(cfg_dir)]
    assert CliRunner().invoke(app, args).exit_code == 0
    rows = _dump(db)
    bad = DayCount("x", date(2019, 1, 2), -1, True)  # violates CHECK (count >= 0) mid-insert
    real = counters.zero_days_invalid
    monkeypatch.setattr(counters, "zero_days_invalid", lambda days: real(days) + [bad])
    assert CliRunner().invoke(app, args).exit_code != 0
    assert _dump(db) == rows


def test_counters_check_flags_a_city_with_no_index_year(tmp_path):
    cfg_dir, data, db = _setup_city(tmp_path, "berlin", BERLIN_ENTRY, [BERLIN_XLSX])
    common = ["--db", str(db), "--config-dir", str(cfg_dir)]
    load = CliRunner().invoke(app, ["counters", "load", "berlin", "--data", str(data), *common])
    assert load.exit_code == 0, load.output
    res = CliRunner().invoke(app, ["counters", "--check", *common])
    assert res.exit_code == 1
    assert "berlin: no index year" in res.output


def _check_db(tmp_path):
    conn = sqlite3.connect(tmp_path / "c.sqlite")
    create_schema(conn)
    return conn


def _site(city, site_id, lat=52.5, lon=13.4):
    return CounterSite(city, site_id, site_id, lat, lon, "s", "l", "u", "t", "2019-01-01", None)


def _cfg(*cities):
    entries = [{**BERLIN_ENTRY, "city": c} for c in cities]
    return CountersConfig.model_validate({"cities": entries, "none": []})


def test_check_passes_a_complete_city(tmp_path):
    conn = _check_db(tmp_path)
    replace_counter_sites(conn, "berlin", [_site("berlin", "a")])
    replace_counter_index(
        conn,
        "berlin",
        [
            CounterIndexYear("berlin", 2019, 100.0, 5, 2019, "counts"),
            CounterIndexYear("berlin", 2020, 110.0, 4, 2019, "counts"),
        ],
    )
    assert check_counters(conn, _cfg("berlin")) == []


def test_check_flags_missing_sites_base_not_100_and_a_site_under_two_cities(tmp_path):
    conn = _check_db(tmp_path)
    replace_counter_sites(conn, "berlin", [_site("berlin", "a")])
    replace_counter_sites(conn, "paris", [_site("paris", "a", 48.8, 2.3)])
    replace_counter_index(
        conn,
        "berlin",
        [
            CounterIndexYear("berlin", 2019, 99.0, 5, 2019, "counts"),
            CounterIndexYear("berlin", 2020, 110.0, 4, 2019, "counts"),
        ],
    )
    problems = check_counters(conn, _cfg("berlin", "paris", "vienna"))
    text = "\n".join(problems)
    assert "berlin: index base year 2019 is 99.0, not 100" in text
    assert "vienna: no counter sites" in text
    assert "paris: no index year" in text
    assert "site 'a' under more than one city: berlin, paris" in text


def test_check_flags_a_step_down_site_absent_from_the_data(tmp_path):
    conn = _check_db(tmp_path)
    replace_counter_sites(conn, "berlin", [_site("berlin", "a")])
    replace_counter_index(
        conn,
        "berlin",
        [
            CounterIndexYear("berlin", 2019, 100.0, 5, 2019, "counts"),
            CounterIndexYear("berlin", 2020, 110.0, 4, 2019, "counts"),
        ],
    )
    entry = {**BERLIN_ENTRY, "step_downs": [{"site": "a", "since": "2020-04"},
                                            {"site": "renamed", "since": "2020-05"}]}  # fmt: skip
    cfg = CountersConfig.model_validate({"cities": [entry], "none": []})
    assert check_counters(conn, cfg) == ["berlin: step_downs site 'renamed' is not a stored site"]


# --- the base year needs MIN_SITES qualifying sites (spec §2.13; follow-up Task 3) -----------


def test_a_base_year_with_fewer_than_min_sites_qualifying_sites_has_no_value():
    """Spec §2.13: a year with fewer than MIN_SITES qualifying sites has no index value. That
    holds for the base year too, and the chain from it then has no value either."""
    few = MIN_SITES - 1
    days = [d for i in range(few) for y in (2022, 2023) for d in _year(f"s{i}", y, 10)]
    assert chain_index(days, 2022) == [(2022, None, few), (2023, None, few)]
    enough = days + [d for y in (2022, 2023) for d in _year("extra", y, 10)]
    assert chain_index(enough, 2022) == [(2022, 100.0, MIN_SITES), (2023, 100.0, MIN_SITES)]


def test_a_monthly_base_year_with_fewer_than_min_sites_gauges_has_no_value():
    def gauge(g, year):
        return [MonthValue(g, f"{year}-{m:02d}", "working_day", 50.0) for m in range(1, 13)]

    few = MIN_SITES - 1
    months = [v for i in range(few) for y in (2019, 2020) for v in gauge(f"g{i}", y)]
    assert chain_index_monthly(months, 2019, "working_day") == [
        (2019, None, few),
        (2020, None, few),
    ]


def test_check_accepts_no_value_on_a_base_year_with_too_few_sites(tmp_path):
    conn = _check_db(tmp_path)
    replace_counter_sites(conn, "berlin", [_site("berlin", "a")])
    few = MIN_SITES - 1
    replace_counter_index(
        conn,
        "berlin",
        [
            CounterIndexYear("berlin", 2019, None, few, 2019, "counts"),
            CounterIndexYear("berlin", 2020, None, few, 2019, "counts"),
        ],
    )
    assert check_counters(conn, _cfg("berlin")) == ["berlin: no index year"]


def test_check_flags_a_value_on_a_base_year_with_too_few_sites(tmp_path):
    conn = _check_db(tmp_path)
    replace_counter_sites(conn, "berlin", [_site("berlin", "a")])
    few = MIN_SITES - 1
    replace_counter_index(
        conn,
        "berlin",
        [
            CounterIndexYear("berlin", 2019, 100.0, few, 2019, "counts"),
            CounterIndexYear("berlin", 2020, 110.0, MIN_SITES, 2019, "counts"),
        ],
    )
    assert check_counters(conn, _cfg("berlin")) == [
        f"berlin: index base year 2019 has {few} qualifying sites (fewer than {MIN_SITES}) "
        "but the value 100.0"
    ]
